"""Export Excel du journal d'audit (architecture § 5.6), avec openpyxl en mode écriture seule.

- Borné : au plus ``ALPHA_AUDIT_EXPORT_MAX_LIGNES`` événements (100 000 par défaut, § 5.6) ;
  au-delà, ``ExportTropVolumineux`` : l'appelant doit restreindre ses filtres.
- Toutes les cellules texte sont écrites en type « chaîne » : une valeur qui commence par
  ``=``, ``+``, ``-`` ou ``@`` (nom d'objet saisi par un utilisateur) n'est jamais une formule.
- Feuille « Journal » (ordre de la chaîne, rang croissant) puis feuille « Export » (qui,
  quand, quels filtres, combien de lignes).
"""

from __future__ import annotations

import datetime as dt
import io
import json
from typing import Any

from django.conf import settings
from django.db.models import QuerySet
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from .libelles import libelle_action
from .models import EvenementAudit

TYPE_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_LIGNES_DEFAUT = 100_000

COLONNES = (
    ("Horodatage", 22),
    ("Rang", 8),
    ("Type d'acteur", 14),
    ("Acteur", 26),
    ("Catégorie", 15),
    ("Action", 28),
    ("Libellé de l'action", 34),
    ("Type d'objet", 24),
    ("Identifiant de l'objet", 38),
    ("Objet", 30),
    ("Détail", 50),
    ("Avant", 40),
    ("Après", 40),
    ("Adresse IP", 16),
    ("Empreinte", 66),
    ("Empreinte précédente", 66),
)
_CHAMPS = (
    "horodatage",
    "seq",
    "acteur_type",
    "acteur_libelle",
    "categorie",
    "action",
    "objet_type",
    "objet_id",
    "objet_libelle",
    "detail",
    "avant",
    "apres",
    "ip",
    "hash",
    "hash_precedent",
)


class ExportTropVolumineux(Exception):
    def __init__(self, maximum: int) -> None:
        super().__init__(f"Plus de {maximum} lignes.")
        self.maximum = maximum


def max_lignes() -> int:
    return int(getattr(settings, "ALPHA_AUDIT_EXPORT_MAX_LIGNES", MAX_LIGNES_DEFAUT))


def _texte(feuille, valeur: Any, *, gras: bool = False) -> WriteOnlyCell:
    cellule = WriteOnlyCell(feuille, value=ILLEGAL_CHARACTERS_RE.sub("", str(valeur)))
    cellule.data_type = "s"  # jamais une formule
    if gras:
        cellule.font = Font(bold=True)
    return cellule


def _json(valeur: Any) -> str:
    return "" if valeur is None else json.dumps(valeur, ensure_ascii=False, sort_keys=True)


def _ligne(feuille, evenement: dict[str, Any]) -> list[Any]:
    quand = timezone.localtime(evenement["horodatage"]).strftime("%Y-%m-%d %H:%M:%S.%f")
    cellules: list[Any] = [_texte(feuille, quand)]
    for champ in _CHAMPS[1:]:
        valeur = evenement[champ]
        if champ == "seq":
            cellules.append(valeur)  # entier
            continue
        if champ in ("avant", "apres"):
            valeur = _json(valeur)
        cellules.append(_texte(feuille, "" if valeur is None else str(valeur).strip()))
    # Le libellé de l'action vient juste après le code de l'action.
    cellules.insert(6, _texte(feuille, libelle_action(evenement["action"])))
    return cellules


def generer(
    evenements: QuerySet[EvenementAudit],
    *,
    etablissement: str,
    exporte_par: str,
    filtres: dict[str, Any],
) -> tuple[bytes, int]:
    """Classeur (octets) et nombre de lignes ; ``ExportTropVolumineux`` au-delà de la borne."""
    maximum = max_lignes()
    # Contrôle AVANT d'ouvrir le classeur (un classeur en écriture seule abandonné en cours
    # de feuille laisse un fichier temporaire à demi écrit).
    if evenements.order_by()[: maximum + 1].count() > maximum:
        raise ExportTropVolumineux(maximum)
    classeur = Workbook(write_only=True)
    try:
        return _generer(classeur, evenements, etablissement, exporte_par, filtres)
    finally:
        classeur.close()


def _generer(
    classeur: Workbook,
    evenements: QuerySet[EvenementAudit],
    etablissement: str,
    exporte_par: str,
    filtres: dict[str, Any],
) -> tuple[bytes, int]:
    maximum = max_lignes()

    maximum = max_lignes()
    feuille = classeur.create_sheet("Journal")
    for i, (_, largeur) in enumerate(COLONNES, start=1):
        feuille.column_dimensions[chr(64 + i)].width = largeur
    feuille.freeze_panes = "A2"
    feuille.append([_texte(feuille, titre, gras=True) for titre, _ in COLONNES])

    lignes = evenements.order_by("seq").values(*_CHAMPS)
    nb = 0
    for evenement in lignes.iterator(chunk_size=2000):
        nb += 1
        if nb > maximum:
            raise ExportTropVolumineux(maximum)
        feuille.append(_ligne(feuille, evenement))

    infos = classeur.create_sheet("Export")
    infos.column_dimensions["A"].width = 26
    infos.column_dimensions["B"].width = 80
    maintenant = timezone.localtime().strftime("%Y-%m-%d %H:%M:%S")
    rubriques = [
        ("Établissement", etablissement),
        ("Exporté par", exporte_par),
        ("Exporté le", maintenant),
        ("Nombre d'événements", nb),
        ("Filtres", _json(filtres) if filtres else "aucun"),
        (
            "Intégrité",
            "Les colonnes Empreinte et Empreinte précédente permettent de revérifier la chaîne.",
        ),
    ]
    for nom, valeur in rubriques:
        infos.append(
            [
                _texte(infos, nom, gras=True),
                valeur if isinstance(valeur, int) else _texte(infos, valeur),
            ]
        )

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue(), nb


def nom_de_fichier(code_etablissement: str, maintenant: dt.datetime | None = None) -> str:
    instant = timezone.localtime(maintenant or timezone.now())
    sur = (
        "".join(c for c in code_etablissement.lower() if c.isalnum() or c in "-_")
        or "etablissement"
    )
    return f"journal-audit-{sur}-{instant:%Y%m%d-%H%M%S}.xlsx"
