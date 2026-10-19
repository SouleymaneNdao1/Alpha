"""API du journal d'audit ``/api/v1/audit`` (tâche B11, § 5.6) : permissions, isolation A/B,
filtres, curseur, indicateurs, export Excel journalisé, contrôle d'intégrité lancé."""

from __future__ import annotations

import datetime as dt
import io
import uuid
from unittest import mock

import pytest
from django.utils import timezone
from drf_spectacular.generators import SchemaGenerator
from drf_spectacular.validation import validate_schema
from openpyxl import load_workbook

from alpha.audit.models import (
    Categorie,
    ControleIntegriteAudit,
    Declencheur,
    EvenementAudit,
    StatutControle,
)
from alpha.audit.services import journaliser
from alpha.organisation.codes_permission import CodePermission as C
from alpha.tenancy.contexte import contexte_tenant

from .fabriques import UtilisateurFactory
from .outils_admin import avec_permissions, evenements

pytestmark = pytest.mark.django_db

API = "/api/v1/audit"
LECTURE, EXPORT = C.AUDIT_LECTURE, C.AUDIT_EXPORT
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: (méthode, chemin, permission requise)
ROUTES = [
    ("get", "/evenements", LECTURE),
    ("get", "/indicateurs", LECTURE),
    ("get", "/export.xlsx", EXPORT),
    ("get", "/controles", LECTURE),
    ("post", "/controles", LECTURE),
]


def ecrire(etablissement, action="test.evenement", categorie=Categorie.MODIFICATION, **champs):
    """Événement dans la chaîne de ``etablissement`` ; ``quand`` fixe l'horodatage."""
    quand = champs.pop("quand", None)
    with contexte_tenant(etablissement.id):
        if quand is None:
            return journaliser(action, categorie, **champs)
        with mock.patch("django.utils.timezone.now", return_value=quand):
            return journaliser(action, categorie, **champs)


@pytest.fixture
def auditeur(banque):
    return avec_permissions(banque, LECTURE, EXPORT, prenom="Awa", nom="Diop")


@pytest.fixture
def client_auditeur(client_connecte, auditeur):
    return client_connecte(auditeur)


def lister(client, **params):
    reponse = client.get(f"{API}/evenements", params)
    assert reponse.status_code == 200, reponse.content
    return reponse.json()


def actions(corps):
    return [e["action"] for e in corps["results"]]


# --- Accès : 401, 403, super-admin ---------------------------------------------------------


@pytest.mark.parametrize(("methode", "chemin", "perm"), ROUTES)
def test_anonyme_401(client_api, methode, chemin, perm):
    reponse = getattr(client_api, methode)(f"{API}{chemin}")
    assert reponse.status_code == 401
    assert reponse.json()["code"] == "NON_AUTHENTIFIE"


@pytest.mark.parametrize(("methode", "chemin", "perm"), ROUTES)
def test_sans_la_permission_403(client_connecte, banque, methode, chemin, perm):
    autres = [c for c in C.values if c != perm]
    client = client_connecte(avec_permissions(banque, *autres))
    reponse = getattr(client, methode)(f"{API}{chemin}")
    assert reponse.status_code == 403
    assert reponse.json()["code"] == "PERMISSION_REFUSEE"


@pytest.mark.parametrize(("methode", "chemin", "perm"), ROUTES)
def test_super_admin_fyxop_403(client_connecte, methode, chemin, perm):
    client = client_connecte(UtilisateurFactory(super_admin=True))
    reponse = getattr(client, methode)(f"{API}{chemin}")
    assert reponse.status_code == 403
    assert reponse.json()["code"] == "PERMISSION_REFUSEE"


@pytest.mark.parametrize(("methode", "chemin", "perm"), ROUTES)
def test_la_permission_ouvre_la_route(client_connecte, banque, methode, chemin, perm):
    client = client_connecte(avec_permissions(banque, perm))
    reponse = getattr(client, methode)(f"{API}{chemin}")
    assert reponse.status_code in (200, 202), reponse.content


def test_lecture_n_ouvre_pas_l_export_et_inversement(client_connecte, banque):
    seulement_lecture = client_connecte(avec_permissions(banque, LECTURE))
    assert seulement_lecture.get(f"{API}/export.xlsx").status_code == 403
    seulement_export = client_connecte(avec_permissions(banque, EXPORT))
    assert seulement_export.get(f"{API}/evenements").status_code == 403


def test_methodes_non_prevues_405(client_auditeur):
    assert client_auditeur.post(f"{API}/evenements").status_code == 405
    assert client_auditeur.delete(f"{API}/controles").status_code == 405
    assert client_auditeur.put(f"{API}/evenements").status_code == 405


# --- Isolation A/B ---------------------------------------------------------------------------


def test_liste_limitee_a_l_etablissement(client_auditeur, banque, autre_banque):
    ecrire(banque, "test.de_a")
    ecrire(autre_banque, "test.de_b", objet_libelle="DOSSIER-SECRET-B")
    corps = lister(client_auditeur)
    assert "test.de_a" in actions(corps)
    assert "test.de_b" not in actions(corps)
    assert "DOSSIER-SECRET-B" not in str(corps)


def test_chaine_plateforme_invisible_depuis_un_etablissement(client_auditeur, banque):
    with contexte_tenant(None, scope="platform"):
        journaliser("test.plateforme", Categorie.SUPER_ADMIN, etablissement_id=None)
    ecrire(banque, "test.de_a")
    assert "test.plateforme" not in actions(lister(client_auditeur))


def test_filtre_sur_un_acteur_de_b_ne_donne_rien(client_auditeur, autre_banque):
    acteur_b = UtilisateurFactory(etablissement=autre_banque)
    ecrire(autre_banque, "test.de_b", acteur=acteur_b)
    assert lister(client_auditeur, acteur_id=str(acteur_b.id))["results"] == []


def test_indicateurs_et_controles_limites_a_l_etablissement(
    client_auditeur, client_connecte, autre_banque
):
    auditeur_b = avec_permissions(autre_banque, LECTURE)
    ecrire(autre_banque, "auth.connexion.echec", Categorie.CONNEXION)
    ecrire(autre_banque, "utilisateur.invite", Categorie.PARAMETRAGE)
    client_b = client_connecte(auditeur_b)
    assert client_b.post(f"{API}/controles").status_code == 202

    ind = client_auditeur.get(f"{API}/indicateurs").json()
    assert ind["connexions_echouees_24h"] == 0
    assert ind["actions_administration_24h"] == 0
    assert ind["dernier_controle"] is None and ind["controle_en_cours"] is False
    assert client_auditeur.get(f"{API}/controles").json()["count"] == 0


# --- Liste : forme, filtres, curseur ----------------------------------------------------------


def test_forme_d_un_evenement(client_auditeur, banque, auditeur):
    ecrire(
        banque,
        "role.modifie",
        Categorie.PARAMETRAGE,
        objet=("organisation.role", "r-1"),
        objet_libelle="DAF",
        avant={"avis": None},
        apres={"avis": "favorable"},
        detail="Avis : — → favorable",
        acteur=auditeur,
    )
    e = lister(client_auditeur, action="role.modifie")["results"][0]
    assert e["action_libelle"] == "Modification d'un rôle"
    assert (e["objet_type"], e["objet_id"], e["objet_libelle"]) == (
        "organisation.role",
        "r-1",
        "DAF",
    )
    assert e["avant"] == {"avis": None} and e["apres"] == {"avis": "favorable"}
    assert e["acteur_libelle"] == "Awa Diop" and e["acteur_type"] == "UTILISATEUR"
    assert len(e["hash"]) == 64 and e["hash_court"] == f"{e['hash'][:4]}…{e['hash'][-4:]}"
    assert {"seq", "horodatage", "categorie", "detail", "ip"} <= set(e)


def test_ordre_decroissant_et_code_inconnu_affiche_tel_quel(client_auditeur, banque):
    for i in range(3):
        ecrire(banque, f"test.ordre{i}")
    resultats = lister(client_auditeur, action="test.ordre*")["results"]
    assert [e["action"] for e in resultats] == ["test.ordre2", "test.ordre1", "test.ordre0"]
    assert resultats[0]["action_libelle"] == "test.ordre2"
    assert [e["seq"] for e in resultats] == sorted((e["seq"] for e in resultats), reverse=True)


def test_curseur_parcourt_tout_sans_doublon(client_auditeur, banque):
    for i in range(7):
        ecrire(banque, f"test.page{i}")
    vus, url, pages = [], f"{API}/evenements?action=test.page*&taille=3", 0
    while url:
        corps = client_auditeur.get(url).json()
        assert set(corps) == {"next", "results"}
        assert len(corps["results"]) <= 3
        vus += [e["seq"] for e in corps["results"]]
        url, pages = corps["next"], pages + 1
    assert pages == 3 and len(vus) == len(set(vus)) == 7
    assert vus == sorted(vus, reverse=True)


def test_curseur_invalide_400(client_auditeur):
    reponse = client_auditeur.get(f"{API}/evenements?curseur=n-importe-quoi")
    assert reponse.status_code == 400
    assert reponse.json()["code"] == "VALIDATION"
    assert "curseur" in reponse.json()["details"]


def test_taille_bornee(client_auditeur, banque):
    for i in range(3):
        ecrire(banque, f"test.t{i}")
    assert len(lister(client_auditeur, action="test.t*", taille=2)["results"]) == 2
    assert len(lister(client_auditeur, action="test.t*", taille=1000)["results"]) == 3


def test_filtre_categorie(client_auditeur, banque):
    ecrire(banque, "test.a", Categorie.CONNEXION)
    ecrire(banque, "test.b", Categorie.DECISION)
    ecrire(banque, "test.c", Categorie.PARAMETRAGE)
    assert actions(lister(client_auditeur, categorie="CONNEXION", action="test.*")) == ["test.a"]
    deux = actions(lister(client_auditeur, categorie="CONNEXION,DECISION", action="test.*"))
    assert sorted(deux) == ["test.a", "test.b"]


def test_filtre_action_exacte_et_prefixe(client_auditeur, banque):
    ecrire(banque, "utilisateur.invite", Categorie.PARAMETRAGE)
    ecrire(banque, "utilisateur.suspendu", Categorie.SECURITE)
    ecrire(banque, "role.cree", Categorie.PARAMETRAGE)
    assert actions(lister(client_auditeur, action="role.cree")) == ["role.cree"]
    assert sorted(actions(lister(client_auditeur, action="utilisateur.*"))) == [
        "utilisateur.invite",
        "utilisateur.suspendu",
    ]
    assert sorted(actions(lister(client_auditeur, action="role.cree,utilisateur.invite"))) == [
        "role.cree",
        "utilisateur.invite",
    ]


def test_filtre_acteur_et_objet(client_auditeur, banque, auditeur):
    autre = UtilisateurFactory(etablissement=banque, prenom="Mamadou", nom="Sow")
    ecrire(banque, "test.par_awa", acteur=auditeur, objet=("credit.dossier", "42"))
    ecrire(banque, "test.par_mamadou", acteur=autre, objet=("credit.dossier", "43"))
    ecrire(banque, "test.autre_objet", acteur=autre, objet=("tiers.client", "42"))
    assert actions(lister(client_auditeur, acteur_id=str(auditeur.id))) == ["test.par_awa"]
    assert sorted(actions(lister(client_auditeur, objet_type="credit.dossier"))) == [
        "test.par_awa",
        "test.par_mamadou",
    ]
    assert sorted(actions(lister(client_auditeur, objet_id="42"))) == [
        "test.autre_objet",
        "test.par_awa",
    ]
    assert actions(lister(client_auditeur, objet_type="credit.dossier", objet_id="43")) == [
        "test.par_mamadou"
    ]


def test_recherche_texte(client_auditeur, banque, auditeur):
    ecrire(banque, "test.q1", acteur=auditeur, objet_libelle="CR-2026-0412")
    ecrire(banque, "test.q2", objet_libelle="Autre", detail="contient le mot PINGOUIN")
    with (
        contexte_tenant(banque.id),
        mock.patch("alpha.core.contexte.ip", return_value="10.20.4.17"),
    ):
        journaliser("test.q3", Categorie.CONNEXION)
    assert actions(lister(client_auditeur, q="cr-2026")) == ["test.q1"]
    assert actions(lister(client_auditeur, q="awa d")) == ["test.q1"]
    assert actions(lister(client_auditeur, q="pingouin")) == ["test.q2"]
    assert actions(lister(client_auditeur, q="10.20.4")) == ["test.q3"]
    assert lister(client_auditeur, q="%")["results"] == []  # caractère joker échappé


def test_filtres_de_periode(client_auditeur, banque):
    maintenant = timezone.now()
    ecrire(banque, "test.maintenant")
    ecrire(banque, "test.il_y_a_3j", quand=maintenant - dt.timedelta(days=3))
    ecrire(banque, "test.il_y_a_20j", quand=maintenant - dt.timedelta(days=20))
    ecrire(banque, "test.il_y_a_60j", quand=maintenant - dt.timedelta(days=60))

    def p(**params):
        return {a for a in actions(lister(client_auditeur, action="test.*", **params))}

    assert p(periode="aujourdhui") == {"test.maintenant"}
    assert p(periode="7j") == {"test.maintenant", "test.il_y_a_3j"}
    assert p(periode="30j") == {"test.maintenant", "test.il_y_a_3j", "test.il_y_a_20j"}
    assert len(p()) == 4


def test_filtres_du_au(client_auditeur, banque):
    ecrire(banque, "test.premier", quand=dt.datetime(2026, 3, 10, 8, 0, tzinfo=dt.UTC))
    ecrire(banque, "test.dernier", quand=dt.datetime(2026, 3, 10, 23, 59, 59, tzinfo=dt.UTC))
    ecrire(banque, "test.lendemain", quand=dt.datetime(2026, 3, 11, 0, 0, tzinfo=dt.UTC))

    def p(**params):
        return {a for a in actions(lister(client_auditeur, action="test.*", **params))}

    # Dakar = UTC : un jour entier, bornes incluses.
    assert p(du="2026-03-10", au="2026-03-10") == {"test.premier", "test.dernier"}
    assert p(du="2026-03-11") == {"test.lendemain"}
    assert p(au="2026-03-10") == {"test.premier", "test.dernier"}
    assert p(du="2026-03-10T09:00:00Z", au="2026-03-10T23:59:59Z") == {"test.dernier"}


@pytest.mark.parametrize(
    ("params", "champ"),
    [
        ({"categorie": "INCONNUE"}, "categorie"),
        ({"periode": "1an"}, "periode"),
        ({"du": "hier"}, "du"),
        ({"du": "2026-03-10", "au": "2026-03-01"}, "du"),
        ({"periode": "7j", "du": "2026-03-01"}, "periode"),
        ({"acteur_id": "pas-un-uuid"}, "acteur_id"),
    ],
)
def test_filtres_invalides_400(client_auditeur, params, champ):
    for chemin in ("evenements", "export.xlsx"):
        reponse = client_auditeur.get(f"{API}/{chemin}", params)
        assert reponse.status_code == 400
        corps = reponse.json()
        assert corps["code"] == "VALIDATION" and champ in corps["details"]


# --- Indicateurs ------------------------------------------------------------------------------


def test_indicateurs(client_auditeur, banque):
    maintenant = timezone.now()
    ecrire(banque, "auth.connexion.echec", Categorie.CONNEXION)
    ecrire(banque, "auth.connexion.echec", Categorie.CONNEXION)
    ecrire(
        banque,
        "auth.connexion.echec",
        Categorie.CONNEXION,
        quand=maintenant - dt.timedelta(hours=30),
    )
    ecrire(banque, "auth.connexion.succes", Categorie.CONNEXION)
    ecrire(banque, "utilisateur.invite", Categorie.PARAMETRAGE)
    ecrire(banque, "role.modifie", Categorie.PARAMETRAGE)
    ecrire(banque, "role.cree", Categorie.PARAMETRAGE, quand=maintenant - dt.timedelta(days=2))

    corps = client_auditeur.get(f"{API}/indicateurs").json()
    assert corps["connexions_echouees_24h"] == 2
    assert corps["actions_administration_24h"] == 2
    assert corps["evenements_aujourdhui"] >= 5  # + les événements de création du jeu de test
    assert corps["dernier_controle"] is None
    assert corps["controle_en_cours"] is False


def test_indicateur_dernier_controle(client_auditeur, banque):
    ecrire(banque, "test.x")
    with contexte_tenant(banque.id):
        ControleIntegriteAudit.objects.create(
            etablissement_id=banque.id,
            lance_le=timezone.now() - dt.timedelta(hours=2),
            termine_le=timezone.now() - dt.timedelta(hours=2),
            declencheur=Declencheur.PLANIFIE,
            statut=StatutControle.ROMPUE,
            dernier_seq=4,
            premier_seq_invalide=3,
            message="Chaîne rompue au seq 3.",
        )
        ControleIntegriteAudit.objects.create(
            etablissement_id=banque.id,
            lance_le=timezone.now(),
            declencheur=Declencheur.MANUEL,
            statut=StatutControle.EN_COURS,
        )
    corps = client_auditeur.get(f"{API}/indicateurs").json()
    assert corps["controle_en_cours"] is True
    assert corps["dernier_controle"]["statut"] == "ROMPUE"  # le contrôle en cours est ignoré
    assert corps["dernier_controle"]["premier_seq_invalide"] == 3
    assert corps["dernier_controle"]["dernier_seq"] == 4


# --- Export Excel -------------------------------------------------------------------------------


def ouvrir_export(reponse):
    assert reponse.status_code == 200, reponse.content
    return load_workbook(io.BytesIO(reponse.content))


def lignes_journal(classeur):
    return [list(r) for r in classeur["Journal"].iter_rows(values_only=True)]


def test_export_contenu(client_auditeur, banque, auditeur):
    for i in range(3):
        ecrire(
            banque,
            f"test.export{i}",
            Categorie.DECISION,
            acteur=auditeur,
            objet=("credit.dossier", str(i)),
            objet_libelle=f"CR-{i}",
            avant={"a": i},
            apres={"a": i + 1},
            detail=f"détail {i} é",
        )
    reponse = client_auditeur.get(f"{API}/export.xlsx", {"action": "test.export*"})
    assert reponse["Content-Type"] == XLSX
    assert "attachment" in reponse["Content-Disposition"]
    assert ".xlsx" in reponse["Content-Disposition"]
    assert reponse["Cache-Control"] == "no-store"
    classeur = ouvrir_export(reponse)
    assert classeur.sheetnames == ["Journal", "Export"]
    lignes = lignes_journal(classeur)
    entete, donnees = lignes[0], lignes[1:]
    assert entete[:3] == ["Horodatage", "Rang", "Type d'acteur"]
    assert len(donnees) == 3
    colonne = {nom: i for i, nom in enumerate(entete)}
    assert [r[colonne["Action"]] for r in donnees] == [
        "test.export0",
        "test.export1",
        "test.export2",
    ]
    premiere = donnees[0]
    assert premiere[colonne["Acteur"]] == "Awa Diop"
    assert premiere[colonne["Catégorie"]] == "DECISION"
    assert premiere[colonne["Objet"]] == "CR-0"
    assert premiere[colonne["Détail"]] == "détail 0 é"
    assert premiere[colonne["Avant"]] == '{"a": 0}' and premiere[colonne["Après"]] == '{"a": 1}'
    assert isinstance(premiere[colonne["Rang"]], int)
    seqs = [r[colonne["Rang"]] for r in donnees]
    assert seqs == sorted(seqs)
    with contexte_tenant(banque.id):
        attendu = EvenementAudit.objects.get(seq=seqs[0])
    assert premiere[colonne["Empreinte"]] == attendu.hash
    assert premiere[colonne["Empreinte précédente"]] == attendu.hash_precedent

    infos = {r[0]: r[1] for r in classeur["Export"].iter_rows(values_only=True)}
    assert infos["Nombre d'événements"] == 3
    assert "Awa Diop" in infos["Exporté par"]
    assert "test.export*" in infos["Filtres"]


def test_export_respecte_les_filtres_et_l_isolation(client_auditeur, banque, autre_banque):
    ecrire(banque, "test.garde", Categorie.DECISION)
    ecrire(banque, "test.ecarte", Categorie.CONNEXION)
    ecrire(autre_banque, "test.garde", Categorie.DECISION, objet_libelle="SECRET-B")
    reponse = client_auditeur.get(
        f"{API}/export.xlsx", {"categorie": "DECISION", "action": "test.*"}
    )
    lignes = lignes_journal(ouvrir_export(reponse))
    assert len(lignes) == 2
    assert "SECRET-B" not in str(lignes)


def test_export_les_formules_ne_s_executent_pas(client_auditeur, banque):
    ecrire(banque, "test.formule", objet_libelle='=HYPERLINK("http://x","clic")', detail="+1+1")
    classeur = ouvrir_export(client_auditeur.get(f"{API}/export.xlsx", {"action": "test.formule"}))
    cellules = [c for ligne in classeur["Journal"].iter_rows(min_row=2) for c in ligne]
    formule = next(c for c in cellules if str(c.value).startswith("=HYPERLINK"))
    assert formule.data_type == "s"
    assert any(c.value == "+1+1" and c.data_type == "s" for c in cellules)


def test_export_est_journalise(client_auditeur, banque, auditeur):
    assert evenements(banque, "audit.export") == []
    ecrire(banque, "test.a_exporter")
    reponse = client_auditeur.get(
        f"{API}/export.xlsx", {"action": "test.*", "categorie": "MODIFICATION"}
    )
    assert reponse.status_code == 200
    (trace,) = evenements(banque, "audit.export")
    assert trace.categorie == Categorie.DOCUMENT
    assert trace.acteur_id == auditeur.id and trace.acteur_libelle == "Awa Diop"
    assert trace.apres["nb_lignes"] == 1
    assert trace.apres["filtres"] == {"action": ["test.*"], "categorie": ["MODIFICATION"]}
    assert "1 ligne" in trace.detail
    # L'export suivant contient la trace du précédent.
    suivant = client_auditeur.get(f"{API}/export.xlsx", {"action": "audit.export"})
    assert len(lignes_journal(ouvrir_export(suivant))) == 2
    assert len(evenements(banque, "audit.export")) == 2
    # Visible dans la liste.
    assert "audit.export" in actions(lister(client_auditeur, action="audit.export"))


def test_export_borne(client_auditeur, banque, settings):
    settings.ALPHA_AUDIT_EXPORT_MAX_LIGNES = 3
    for i in range(4):
        ecrire(banque, f"test.gros{i}")
    refus = client_auditeur.get(f"{API}/export.xlsx", {"action": "test.gros*"})
    assert refus.status_code == 400
    corps = refus.json()
    assert corps["code"] == "EXPORT_TROP_VOLUMINEUX" and corps["details"] == {"maximum": 3}
    assert evenements(banque, "audit.export") == []  # un export refusé n'est pas un export
    ok = client_auditeur.get(f"{API}/export.xlsx", {"action": "test.gros0"})
    assert len(lignes_journal(ouvrir_export(ok))) == 2


def test_export_accepte_l_en_tete_accept_du_tableur(client_auditeur, banque):
    ecrire(banque, "test.accept")
    reponse = client_auditeur.get(f"{API}/export.xlsx", {"action": "test.accept"}, HTTP_ACCEPT=XLSX)
    assert reponse.status_code == 200 and reponse["Content-Type"] == XLSX


# --- Contrôle d'intégrité ------------------------------------------------------------------------


def test_lancement_controle_chaine_intacte(
    client_auditeur, banque, auditeur, django_capture_on_commit_callbacks
):
    for i in range(5):
        ecrire(banque, f"test.chaine{i}")
    with django_capture_on_commit_callbacks(execute=True):
        reponse = client_auditeur.post(f"{API}/controles")
    assert reponse.status_code == 202
    corps = reponse.json()
    assert corps["statut"] == "EN_COURS"
    identifiant = uuid.UUID(corps["id"])

    with contexte_tenant(banque.id):
        controle = ControleIntegriteAudit.objects.get(pk=identifiant)
        dernier_seq = EvenementAudit.objects.order_by("-seq").values_list("seq", flat=True).first()
    assert controle.statut == StatutControle.INTACTE
    assert controle.declencheur == Declencheur.MANUEL
    assert controle.termine_le is not None and controle.nb_evenements >= 5
    assert controle.dernier_seq <= dernier_seq
    assert controle.premier_seq_invalide is None

    liste = client_auditeur.get(f"{API}/controles").json()
    assert liste["count"] == 1
    assert liste["results"][0]["id"] == str(identifiant)
    assert liste["results"][0]["statut"] == "INTACTE"
    assert "dernier_hash" not in liste["results"][0]

    ind = client_auditeur.get(f"{API}/indicateurs").json()
    assert ind["dernier_controle"]["statut"] == "INTACTE" and ind["controle_en_cours"] is False

    (trace,) = evenements(banque, "audit.controle")
    assert trace.acteur_id == auditeur.id and trace.objet_id == str(identifiant)


def test_lancement_controle_cree_la_ligne_avant_execution(client_auditeur, banque):
    """Sans exécution des callbacks (worker pas encore passé) : contrôle EN_COURS."""
    reponse = client_auditeur.post(f"{API}/controles")
    assert reponse.status_code == 202
    with contexte_tenant(banque.id):
        controle = ControleIntegriteAudit.objects.get(pk=reponse.json()["id"])
    assert controle.statut == StatutControle.EN_COURS and controle.termine_le is None
    assert client_auditeur.get(f"{API}/indicateurs").json()["controle_en_cours"] is True


def test_un_seul_controle_a_la_fois_409(client_auditeur):
    premier = client_auditeur.post(f"{API}/controles")
    second = client_auditeur.post(f"{API}/controles")
    assert second.status_code == 409
    corps = second.json()
    assert corps["code"] == "CONTROLE_EN_COURS"
    assert corps["details"]["id"] == premier.json()["id"]


def test_controle_abandonne_ne_bloque_pas(client_auditeur, banque):
    with contexte_tenant(banque.id):
        ControleIntegriteAudit.objects.create(
            etablissement_id=banque.id,
            lance_le=timezone.now() - dt.timedelta(hours=3),
            declencheur=Declencheur.MANUEL,
            statut=StatutControle.EN_COURS,
        )
    assert client_auditeur.post(f"{API}/controles").status_code == 202


def test_tache_executer_controle_ignore_un_controle_d_un_autre_etablissement(banque, autre_banque):
    from alpha.audit import tasks

    with contexte_tenant(autre_banque.id):
        controle_b = ControleIntegriteAudit.objects.create(
            etablissement_id=autre_banque.id,
            lance_le=timezone.now(),
            declencheur=Declencheur.MANUEL,
            statut=StatutControle.EN_COURS,
        )
    assert tasks.executer_controle.delay(str(banque.id), str(controle_b.id)).get() == "IGNORE"
    with contexte_tenant(autre_banque.id):
        controle_b.refresh_from_db()
    assert controle_b.statut == StatutControle.EN_COURS


# --- Schéma OpenAPI ------------------------------------------------------------


def test_schema_openapi_valide():
    schema = SchemaGenerator().get_schema(request=None, public=True)
    validate_schema(schema)
    chemins = schema["paths"]
    for methode, chemin, _ in ROUTES:
        operation = chemins[f"/api/v1/audit{chemin}"][methode]
        assert {"401", "403"} <= set(operation["responses"]), (methode, chemin)
    liste = chemins["/api/v1/audit/evenements"]["get"]
    noms = {p["name"] for p in liste["parameters"]}
    assert {
        "q",
        "categorie",
        "action",
        "periode",
        "du",
        "au",
        "acteur_id",
        "objet_type",
        "objet_id",
        "curseur",
        "taille",
    } <= noms
    assert XLSX in chemins["/api/v1/audit/export.xlsx"]["get"]["responses"]["200"]["content"]
    assert "202" in chemins["/api/v1/audit/controles"]["post"]["responses"]
